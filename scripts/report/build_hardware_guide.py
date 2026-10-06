# build_hardware_guide.py - makes the self-contained word guide for building the esp32 prototype (to share with a friend)
"""Run:  python scripts/report/build_hardware_guide.py   ->  docs/ESP32_Prototype_Hardware_Guide.docx"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "ESP32_Prototype_Hardware_Guide.docx"
d = Document()
sec = d.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.0)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(4)
for n, sz in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    h = d.styles[n]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    rf = h.element.get_or_add_rPr().find(qn("w:rFonts"))
    for a in ("w:asciiTheme", "w:hAnsiTheme"):
        rf.attrib.pop(qn(a), None)
    rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)


def p(text, bold=False, italic=False, size=None):
    para = d.add_paragraph(); r = para.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return para


def bullets(items, style="List Bullet"):
    for i, it in enumerate(items, start=1):
        if style == "List Number":
            para = d.add_paragraph(f"{i}.  {it}")
            para.paragraph_format.left_indent = Cm(0.9); para.paragraph_format.first_line_indent = Cm(-0.6)
        else:
            d.add_paragraph(it, style=style)


def table(header, rows, widths):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10); shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(10)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    d.add_paragraph()


def code(text, size=8):
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; shade(c, "F2F2F2"); c.text = ""
    first = True
    for ln in text.rstrip("\n").split("\n"):
        para = c.paragraphs[0] if first else c.add_paragraph(); first = False
        para.paragraph_format.space_after = Pt(0)
        r = para.add_run(ln if ln else " "); r.font.name = "Consolas"; r.font.size = Pt(size)
        r._r.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:cs"), "Consolas")
    d.add_paragraph()


t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Smart Waste Sorter Prototype"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Hardware build guide: ESP32 + DHT22 + MQ-135 + servo"); r.font.size = Pt(14)
p("What this builds: a small bench unit that senses humidity and gas next to an item, and tilts a tray left or right (or signals “ask a person”) according to commands from a laptop that "
  "classifies the item from a camera photo. This guide covers the hardware and the ESP32 code, which is what you need to build and test the unit. The laptop software is described in section 9.", italic=True)

d.add_heading("1. Parts list", 1)
table(["Part", "Qty", "Note"], [
    ("ESP32 DevKit (30 or 38 pin) + USB data cable", "1", "must be a data cable, not charge-only"),
    ("DHT22 (AM2302) humidity and temperature sensor", "1", "module with 3 pins, or bare 4-pin sensor + 10 kΩ resistor"),
    ("MQ-135 gas sensor module", "1", "needs 5 V; heater warms up"),
    ("SG90 micro servo", "1", "tilts the tray"),
    ("10 kΩ resistors", "2 (+1 for a bare DHT22)", "voltage divider for the MQ-135"),
    ("Breadboard and jumper wires (M-M and M-F)", "1 set", ""),
    ("Cardboard, tape or hot glue, white paper, 2 cups or small boxes", "–", "tray, frame and bins"),
    ("Optional: 3 LEDs (green, red, yellow) + 3 × 220 Ω, active buzzer, 470 µF capacitor", "–", "indicators; the capacitor stops the ESP32 resetting when the servo moves"),
    ("Camera: a phone with a free IP-webcam app, or a USB webcam", "1", "used by the laptop software"),
], [8.0, 2.8, 5.6])

d.add_heading("2. Wiring", 1)
p("Power the ESP32 from the laptop USB port only. Join all grounds (GND).", bold=True)
table(["Part", "Part pin", "ESP32 pin", "Note"], [
    ("DHT22", "VCC / +", "3V3", ""),
    ("", "GND / −", "GND", ""),
    ("", "DATA / OUT", "GPIO4", "bare 4-pin sensor: 10 kΩ between DATA and 3V3 (seen from the front: pin 1 VCC, 2 DATA, 3 unused, 4 GND)"),
    ("MQ-135", "VCC", "VIN (5 V pin)", "the heater needs 5 V"),
    ("", "GND", "GND", ""),
    ("", "AO", "10 kΩ to GPIO34; second 10 kΩ from GPIO34 to GND", "voltage divider, see warning"),
    ("Servo SG90", "red", "VIN (5 V)", ""),
    ("", "brown or black", "GND", ""),
    ("", "orange or yellow (signal)", "GPIO18", ""),
    ("Trigger", "BOOT button on the board", "GPIO0", "no wiring; press after start-up to classify an item"),
    ("Optional LEDs", "green / red / yellow", "GPIO25 / 26 / 27", "each through 220 Ω to GND"),
    ("Optional buzzer", "+ / −", "GPIO14 / GND", "active buzzer"),
], [2.6, 3.8, 4.6, 5.4])
p("WARNING: never connect the MQ-135 AO pin straight to the ESP32. AO can reach 5 V and the ESP32 pin tolerates only 3.3 V. Use the two 10 kΩ resistors as a divider. "
  "(If you use 10 kΩ + 20 kΩ instead, put the 10 kΩ on the AO side and the 20 kΩ to GND, and change GAS_DIVIDER to 1.5 in the code.)", bold=True)
d.add_picture(str(ROOT / "docs/report/figures/esp32_wiring.png"), width=Cm(16.0))
cap = d.add_paragraph("Figure 1. Wiring diagram", style=None); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].italic = True

d.add_heading("3. Install the tools (about 15 minutes, once)", 1)
bullets([
    "Install Arduino IDE 2.x from arduino.cc.",
    "Plug in the ESP32. If Device Manager shows no COM port, install the USB driver for the chip printed on the board: CP2102 → Silicon Labs CP210x, CH340 → WCH CH340.",
    "File → Preferences → “Additional boards manager URLs”: paste  https://espressif.github.io/arduino-esp32/package_esp32_index.json  → OK.",
    "Tools → Board → Boards Manager → search “esp32” → install “esp32 by Espressif Systems”.",
    "Tools → Manage Libraries → install: ESP32Servo, DHT sensor library (by Adafruit), and Adafruit Unified Sensor (accept when asked).",
], "List Number")

d.add_heading("4. Upload the code", 1)
bullets([
    "Create a new sketch, paste the code from section 10, and save it (for example sorter_node_esp32).",
    "Tools → Board → ESP32 Dev Module. Tools → Port → your COM port.",
    "Click Upload. If it stays on “Connecting…”, hold the BOOT button on the board until the upload starts, then release.",
    "Open Tools → Serial Monitor, 115200 baud, line ending “Newline”. You should see READY.",
], "List Number")

d.add_heading("5. Test the hardware from the Serial Monitor", 1)
table(["Type", "Expected answer", "If it fails"], [
    ("P", "PONG", "wrong port or baud rate; press the EN button once"),
    ("S", '{"m":450,"g":210,"t":29.5,"h":45.0}', "see the checks below"),
    ("A60, A90, A120", "OK and the servo moves", "check signal on GPIO18 and servo power on VIN"),
    ("R", "tray tilts left, then OK", ""),
    ("H", "long beep, tray tilts right, then OK", ""),
    ("V / C", "3 short beeps / 2 long beeps, then OK", "only if the optional LEDs or buzzer are fitted"),
], [3.2, 6.4, 6.8])
bullets([
    "m is the DHT22 humidity × 10 (450 = 45%). m = 0 means the DHT22 gave no reading: check DATA on GPIO4 and the pull-up resistor, and wait 3 seconds after start-up.",
    "g is the MQ-135 reading on a 0–1023 scale. It should rise when you hold a spirit swab, vinegar or something smelly near the sensor. If it is stuck at 0 or 1023, check the voltage divider.",
    "Tune the tray: use A60, A90, A120 to find angles where the tray tilts about 25–30° each way, then change ANGLE_LEFT and ANGLE_RIGHT in the code and upload again.",
    "Close the Serial Monitor before running any Python program: it holds the COM port.",
    "The MQ-135 needs warm-up: keep it powered for at least 30 minutes (24 hours is better) before trusting its numbers.",
])

d.add_heading("6. Build the frame (1–2 hours)", 1)
bullets([
    "Base: a shoebox. Fix the servo on a small bridge in the middle with its axis running front to back, so the tray tilts left and right.",
    "Tray: stiff cardboard about 12 × 14 cm, covered with white paper and clear tape so items slide. Glue a 1.5 cm lip along the front and back edges only. Glue the tray to the servo horn.",
    "Bins: a cup under the left edge (auto-sorted) and one under the right edge (hazardous, mark it red).",
    "Sensor hood: a small cardboard tunnel on one side of the tray with the DHT22 and the MQ-135 inside it, 3–4 cm from where the item sits. The DHT22 senses the air near the item, so it reacts slowly: wait about 10 seconds after placing an item.",
    "Camera: the phone on a stack of books, 25–30 cm straight above the tray; a lamp from the side; plain white background.",
    "Props: light and safe only (under 150 g): plastic bottle, paper cup, cardboard, can, cloth, a dead AA battery, a broken earphone. Never use damaged lithium cells.",
])

d.add_heading("7. Power problems and fixes", 1)
table(["Problem", "Fix"], [
    ("ESP32 resets when the servo moves", "add a 470 µF capacitor across VIN and GND near the servo, or power the servo from a separate 5 V supply with a common ground"),
    ("No COM port", "use a data cable, install the CP210x or CH340 driver, try another USB port"),
    ("Gas reading drifts", "longer warm-up, shield from draughts, keep the room still"),
    ("Item does not slide off the tray", "smoother tray surface, larger tilt angle, lighter item"),
], [6.0, 10.4])

d.add_heading("8. Safety", 1)
p("USB power only. Keep liquids away from the ESP32. Use dead or small cells only, never open or damage a battery.")

d.add_heading("9. Laptop software (for the person who has the project folder)", 1)
p("The ESP32 only senses and moves the tray. The classification and decisions run on the laptop from the project repository (edge-waste). With the project folder:")
code("""pip install pyserial
python scripts/check_node.py --port COM5              # live sensor readings
python scripts/check_node.py --port COM5 --angles     # tray angles
python scripts/check_node.py --port COM5 --actions    # run R, H, V, C
python scripts/prototype_sorter.py --mock-node --image <photo.jpg>     # whole software path, no hardware
python scripts/calibrate_prototype_oci.py collect --port COM5          # then:  ... fit
python scripts/prototype_sorter.py --port COM5 --camera 0              # live: press BOOT or SPACE to classify""", 9)
p("Phone camera: install “IP Webcam” (Android), start the server and use  --camera http://<phone-ip>:8080/video  (same Wi-Fi). Replace COM5 with your port name from Device Manager.")
p("Serial protocol (115200 baud, one command per line): P → PONG; S → JSON with m (humidity × 10), g (gas 0–1023), t, h; R, H, V, C → run an action and answer OK; A<deg> → move the tray; the BOOT button prints BTN.", italic=True)

d.add_heading("10. ESP32 code (complete)", 1)
p("Copy everything in the box into a new Arduino sketch.")
code((ROOT / "hardware/esp32/sorter_node_esp32/sorter_node_esp32.ino").read_text(encoding="utf-8"), 7.5)
p("Note: the code has not been compiled on the author's machine. If the Arduino IDE reports a compile error, send the exact message.", italic=True)

d.save(str(OUT))
print("saved", OUT)
