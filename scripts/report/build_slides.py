# build_slides.py - builds the 25-slide review presentation (pptx) from the same figures and numbers as the report
"""Run:  python scripts/report/build_slides.py     ->  docs/report/BITE497J_Project_I_Review_Slides.pptx"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
REP = ROOT / "docs/report/figures"
SM = ROOT / "docs/report/figures_small"
ML = ROOT / "reports/ml_analysis/figures"
SIM = ROOT / "reports/simulations/figures"
GAL = ROOT / "reports/demo/gallery"
OUT = ROOT / "docs/report/BITE497J_Project_I_Review_Slides.pptx"

NAVY = RGBColor(0x1F, 0x38, 0x64); BLUE = RGBColor(0x2E, 0x75, 0xB6); GREY = RGBColor(0x40, 0x40, 0x40); LIGHT = RGBColor(0xEA, 0xF1, 0xFB)
RED = RGBColor(0xC0, 0x39, 0x2B); GREEN = RGBColor(0x2E, 0x8B, 0x57); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
TOTAL = 25

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
STUDENTS = "Ansh Jerath (23BIT0123)  ·  Arnav Mishra (23BIT0142)  ·  Bhavesh Singh Thakur (23BIT0199)"
SHORT = "Edge-Deployed Waste Classification with Sensor-Augmented Vision  |  BITE497J Project I"


def tb(slide, x, y, w, h, text="", size=16, bold=False, color=GREY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font; r.font.italic = italic
    return box


def bullets(slide, x, y, w, h, items, size=17, gap=5, color=GREY):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        lvl = 0
        if isinstance(it, tuple):
            lvl, it = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        mark = "▪ " if lvl == 0 else "– "
        r = p.add_run(); r.text = ("      " * lvl) + mark + it
        r.font.size = Pt(size - 2 * lvl); r.font.color.rgb = color; r.font.name = FONT
    return box


def picture(slide, path, x, y, w, h):
    """Place the picture inside the box (x, y, w, h) keeping its aspect ratio, centred."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def card(slide, x, y, w, h, title, body, tcolor=BLUE, size=14):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = LIGHT; shp.line.color.rgb = tcolor; shp.line.width = Pt(1.25)
    shp.adjustments[0] = 0.06
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title; r.font.bold = True; r.font.size = Pt(size + 2); r.font.color.rgb = tcolor; r.font.name = FONT
    for line in body:
        q = tf.add_paragraph(); q.space_before = Pt(3)
        rr = q.add_run(); rr.text = line; rr.font.size = Pt(size); rr.font.color.rgb = GREY; rr.font.name = FONT
    return shp


def table(slide, x, y, w, rows, col_w, size=12, header_fill=NAVY, row_h=0.36):
    nr, nc = len(rows), len(rows[0])
    shape = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = shape.table
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(w * cw / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.025)
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val)
            r.font.size = Pt(size); r.font.name = FONT
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            if i == 0:
                c.fill.solid(); c.fill.fore_color.rgb = header_fill; r.font.bold = True; r.font.color.rgb = WHITE
            else:
                c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE; r.font.color.rgb = GREY
    return shape


def base(n, title, notes="", subtitle=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.95))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb(s, 0.45, 0.12, 12.4, 0.75, title, size=28, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.95), prs.slide_width, Inches(0.06))
    acc.fill.solid(); acc.fill.fore_color.rgb = BLUE; acc.line.fill.background()
    tb(s, 0.45, 7.08, 10.5, 0.3, SHORT, size=10, color=RGBColor(0x80, 0x80, 0x80))
    tb(s, 11.6, 7.08, 1.3, 0.3, f"{n} / {TOTAL}", size=10, color=RGBColor(0x80, 0x80, 0x80), align=PP_ALIGN.RIGHT)
    if subtitle:
        tb(s, 0.45, 1.12, 12.4, 0.4, subtitle, size=15, italic=True, color=BLUE)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


# ---------------------------------------------------------------------------------------------- slides
def slide1():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.55), prs.slide_width, Inches(0.07)); stripe.fill.solid(); stripe.fill.fore_color.rgb = BLUE; stripe.line.fill.background()
    tb(s, 0.8, 1.0, 11.7, 1.9, "Edge-Deployed Waste Classification with Sensor-Augmented Vision", size=42, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 0.8, 2.95, 11.7, 0.6, "CAPS-FCL: Contamination-Aware, Privacy-Preserving Smart Federated Classification and Learning", size=20, color=RGBColor(0xBD, 0xD7, 0xEE), italic=True)
    tb(s, 0.8, 3.6, 11.7, 0.5, "BITE497J – Project I  ·  Pre-final (guide) review", size=20, bold=True, color=WHITE)
    tb(s, 0.8, 4.85, 11.7, 0.5, STUDENTS, size=20, color=WHITE)
    tb(s, 0.8, 5.45, 11.7, 0.5, "Guide: Dr. Valarmathi B., Professor Grade 1", size=20, color=WHITE)
    tb(s, 0.8, 6.15, 11.7, 0.5, "School of Computer Science Engineering and Information Systems, VIT, Vellore", size=16, color=RGBColor(0xBD, 0xD7, 0xEE))
    s.notes_slide.notes_text_frame.text = "Introduce the team, the guide and the project title. The project builds an edge-ready waste sorter in software: vision, uncertainty, hazard safety, contamination sensing, explanations and federated learning."


def slide2():
    s = base(2, "Outline", "Walk through the structure: problem, survey, design, results, demonstration, cost and conclusion.")
    left = ["1.  Problem, motivation and objectives", "2.  Literature and patent survey, research gap", "3.  System design and data", "4.  Classifier, training and detector",
            "5.  Safety mechanisms: uncertainty, conformal routing, contamination index"]
    right = ["6.  Results and analysis", "7.  Robustness and simulations", "8.  Demonstration on real images and video", "9.  Cost, schedule, limitations and future work", "10. Conclusion"]
    bullets(s, 0.8, 1.7, 5.9, 4.6, left, size=22, gap=14)
    bullets(s, 7.0, 1.7, 5.9, 4.6, right, size=22, gap=14)


def slide3():
    s = base(3, "Problem and Motivation", "Stress: contamination, hazards, look-alike items, privacy and latency. Figures: World Bank 2018 and the UN Global E-waste Monitor 2024.")
    bullets(s, 0.5, 1.3, 6.4, 5.5, [
        "World waste: 2.01 billion tonnes (2016) rising to 3.40 billion tonnes a year by 2050; at least a third not managed safely",
        "Electronic waste: 62 million tonnes in 2022, rising five times faster than documented recycling",
        "Vision-only sorters give one answer from one model, with few broad categories",
        "They cannot see contamination (grease, food residue)",
        "A wrong hazard call (battery, medical sharp) is far costlier than a wrong plastic bottle",
        "Look-alike items (cardboard boxes vs packaging, steel vs aluminium cans)",
        "Edge latency and privacy: share parameters, not images",
    ], size=17, gap=7)
    picture(s, REP / "fig_lookalike.png", 7.1, 1.45, 5.9, 4.4)
    tb(s, 7.1, 5.95, 5.9, 0.6, "Look-alike classes from the test set: boxes vs packaging (top), steel vs aluminium cans (bottom)", size=12, italic=True, align=PP_ALIGN.CENTER)


def slide4():
    s = base(4, "Problem Statement and Objectives", "Read the problem statement, then the six objectives. Each one has a measurable target on the next slide.")
    card(s, 0.5, 1.3, 12.3, 1.35, "Problem", ["Given an item image (and optional moisture and gas readings), output the item class and material family, a measure of confidence, a contamination "
                                               "score and one action, with a stated limit on hazards reaching recycling, within 500 ms on an edge-class computer."], size=15)
    bullets(s, 0.6, 2.9, 12.2, 4.0, [
        "1. Detect-then-classify: YOLO detector + hybrid ConvNeXt / Vision Transformer, 33 classes in 9 material families",
        "2. Report item accuracy, family (routing) accuracy and hazard recall separately",
        "3. Estimate uncertainty; cross-check object vs material without weakening a hazard; stated hazard-leakage limit",
        "4. Two-sensor Organic Contamination Index that survives a sensor failure",
        "5. Explain decisions: Grad-CAM, LIME, SHAP, Segment Anything outline",
        "6. Federated learning without sharing images; video surveys with one entry per item",
    ], size=17, gap=8)


def slide5():
    s = base(5, "Scope, Assumptions and Measurable Goals", "Software stage on real held-out data; sensors are simulated until the prototype is built.")
    rows = [["Goal", "Measurable target", "Status"],
            ["Fine-grained recognition + routing check", "≥ 15 classes; item / family / hazard recall", "33 classes; 89.30 / 94.14 / 93.50%  (v2: 93.83 / 97.28 / 98.37%)"],
            ["Hazard safety", "Leakage below the operator limit αH", "6.5% → 1.57% at 7.8% review"],
            ["Calibrated uncertainty", "ECE reduced, predictions unchanged", "16.6% → 7.3%"],
            ["Contamination survives sensor loss", "Sensitivity ≥ 95% in each case", "94.9% / 94.5% (simulated)"],
            ["Privacy-preserving learning", "Match a central model", "89.4% vs 89.0% (simulated)"],
            ["Edge readiness", "< 500 ms per item; ONNX parity", "14 ms GPU, 116 ms CPU; 9.4e-6 difference"]]
    table(s, 0.5, 1.35, 12.3, rows, [3.2, 3.6, 4.6], size=13, row_h=0.55)
    bullets(s, 0.6, 5.4, 12.2, 1.6, ["In scope: data, models, decision logic, calibration, explanations, federated simulation, ONNX export, video survey",
                                      "Out of scope for now: physical sensors, conveyor, Raspberry Pi deployment (planned for Project II)"], size=15)


def slide6():
    s = base(6, "Literature Survey: 19 Studies in Five Groups", "Full table with merits and demerits is Table 2.2 of the report. Key message: nobody gives a stated hazard limit, and sensors are rarely fused as a calibrated score.")
    rows = [["Group", "Representative studies", "Typical limitation"],
            ["Vision classification", "Alkılınç 2025 (ensemble), Verber 2026 (hybrid), Nahiduzzaman 2025, Dipo 2025 (YOLOv12), Partosan 2026 (Pi 5)", "Small controlled data; few broad classes; accuracy drops with finer classes"],
            ["Sensors, IoT and fusion", "Chahine 2017, Alnanih 2025, Arun 2025, Hussain 2024, Casao 2024, Chu 2018", "Sensors monitor bins or feed one classifier; no calibrated contamination score"],
            ["Uncertainty", "He 2026 (survey), Radchenko 2024", "Out-of-distribution robustness weak; no hazard limit"],
            ["Federated / edge learning", "Alatawi 2025, McMahan 2017", "Smart bins rarely learn together"],
            ["Datasets and detection", "TACO 2020, TrashNet 2016, SpectralWaste 2024", "Object labels, small or plain-background sets"]]
    table(s, 0.5, 1.35, 12.3, rows, [2.4, 5.4, 4.6], size=13, row_h=0.8)
    tb(s, 0.5, 6.35, 12.3, 0.6, "Findings: single-modal vision dominates · sensors rarely detect residue · ensembles are too heavy for the edge · no stated hazard limit · little private collective learning", size=14, bold=True, color=NAVY)


def slide7():
    s = base(7, "Patent Survey and Research Gap", "Nine related patents were read; the closest multi-sensor reference is CSIRO WO 2024/207048. Table 2.4 compares capabilities.")
    rows = [["Capability", "Vision-only sorters", "Camera + sensor hybrids", "This project"],
            ["More than 15 fine-grained classes", "rare", "not the aim", "33 classes, 9 families"],
            ["Item / family / hazard-recall reported separately", "no", "no", "yes"],
            ["Hazard-exempt object–material cross-check", "no", "no", "yes"],
            ["Calibrated uncertainty + two-tier hazard gate", "rare", "rare", "yes"],
            ["Stated limit on hazards reaching recycling", "no", "no", "yes (conformal sets)"],
            ["Contamination score that survives sensor loss", "no sensors", "extra inputs", "three-model index"],
            ["Explanations; private learning; video inventory", "partly", "rare", "yes; simulated; yes"]]
    table(s, 0.5, 1.35, 12.3, rows, [4.6, 2.4, 2.6, 2.7], size=13, row_h=0.52)
    tb(s, 0.5, 5.75, 12.3, 1.1, "Patents examined: AMP Robotics, Tomra Sorting, Battelle/Sortera, Heil, Fidelity AG, CSIRO, CleanRobotics, PMBFU and an Australian innovation patent. "
       "Differences are listed in Table 2.3 of the report.", size=13, italic=True)


def slide8():
    s = base(8, "Proposed System: Architecture", "Follow the arrows: camera, detector and classifier, cross-check, uncertainty and family sets, sensors, decision engine, actuator.")
    picture(s, REP / "diagram_system.png", 0.5, 1.2, 7.4, 5.8)
    bullets(s, 8.1, 1.4, 4.9, 5.4, [
        "Detector finds the item and names the object",
        "Hybrid classifier gives 33 class probabilities",
        "Cross-check narrows the material, never weakens a hazard",
        "Uncertainty and calibrated family sets bound hazard risk",
        "Moisture + gas sensors give a contamination index",
        "Decision engine: hazardous bin · priority review · contamination reject · family bin",
    ], size=16, gap=8)


def slide9():
    s = base(9, "Data: 33 Classes in 9 Material Families", "Three public Kaggle datasets; coarse folders dropped rather than guessed; TrashBox only for e-waste and medical.")
    picture(s, REP / "fig_taxonomy.png", 0.4, 1.2, 8.3, 5.6)
    bullets(s, 8.9, 1.4, 4.2, 5.4, [
        "28,203 images from 3 Kaggle datasets",
        "27 classes ≈ 500 images; 6 larger merged classes",
        "Largest : smallest class = 11.6",
        "Family = bin; item errors inside a family do not change the route",
        "Hazardous family: battery, e-waste, medical",
        "Split 70 / 15 / 15: 19,739 / 4,232 / 4,232",
    ], size=15, gap=7)


def slide10():
    s = base(10, "Data Pipeline and Leakage Prevention", "Two data bugs were found because results looked too good: path-hash leak and stale duplicates. Both fixed and guarded.")
    picture(s, REP / "diag_data_pipeline.png", 0.4, 1.3, 12.5, 2.6)
    card(s, 0.5, 4.2, 6.0, 2.7, "Bug 1: cross-machine split leak", ["File-name hash used the OS path, so Windows and Colab made different splits", "Re-score gave 94.52% (inflated): ~70% of 'test' images were training images",
                                                                       "Fixed with POSIX relative paths; Colab split recovered: 3,779 / 4,232 = 89.30% exactly"], tcolor=RED, size=13)
    card(s, 6.8, 4.2, 6.0, 2.7, "Bug 2: stale duplicate files", ["Colab showed 1,234 steps per epoch instead of 617", "Old copies of every image on the drive could put one photo in train and test",
                                                                  "Split now ignores files not in the latest provenance record; notebook stops if sizes differ; evaluation refuses a split that does not reproduce the recorded validation accuracy"], tcolor=RED, size=13)


def slide11():
    s = base(11, "Hybrid Classifier: ConvNeXt + Vision Transformer with Attention Fusion", "Two streams: convolution for texture, attention for shape. A learned softmax weight decides per image how much to trust each stream.")
    picture(s, REP / "diag_hybrid_model.png", 0.4, 1.15, 4.0, 5.8)
    bullets(s, 4.7, 1.3, 8.3, 3.2, [
        "ConvNeXt-Tiny 27.8 M + ViT-Small/16 21.7 M; total 50.75 M parameters",
        "GELU activations, LayerNorm (same behaviour at batch 1 on the edge)",
        "Attention fusion over the two streams (measured: ConvNeXt carries 82–93%)",
        "Head: dropout 0.2 → 1024→512 → GELU → dropout 0.2 → 33",
    ], size=16, gap=6)
    box = card(s, 4.7, 4.3, 8.3, 2.5, "Key equations", ["e_s = W2 · GELU(W1 f_s + b1) + b2", "α_s = exp(e_s) / Σ_s' exp(e_s')", "h = [ α_cnx f_cnx ; α_vit f_vit ]   (1,024 values)",
                                                         "p_i = exp(z_i) / Σ_j exp(z_j)     L = −Σ q_i log p_i,  q_i = (1−ε)·1[i=y] + ε/C,  ε = 0.1"], size=15)


def slide12():
    s = base(12, "Training Recipe and the Improved Recipe", "The frozen-feature SVM beat the fine-tuned model, which exposed feature distortion. Training the head first and a 10x lower backbone learning rate fixed it.")
    rows = [["", "First model (v1)", "Improved recipe (v2)"],
            ["Optimiser", "AdamW, lr 3e-4, wd 0.05, one-cycle", "same"],
            ["Order", "all layers from step 1", "backbones frozen for 2 epochs, then fine-tuned"],
            ["Learning rate", "3e-4 everywhere", "3e-5 backbones, 3e-4 new layers"],
            ["Imbalance", "sampler + loss weights (corrects twice)", "sampler only"],
            ["Result: item / family / hazard", "89.30 / 94.14 / 93.50%", "93.83 / 97.28 / 98.37%"]]
    table(s, 0.5, 1.4, 7.4, rows, [2.1, 2.9, 3.0], size=12, row_h=0.62)
    picture(s, REP / "fig_versions.png", 8.1, 1.3, 4.9, 3.0)
    picture(s, SIM / "fig_training_history_v2.png", 0.5, 5.35, 7.4, 1.65)
    bullets(s, 8.1, 4.5, 4.9, 2.4, ["Val accuracy 90.0% after epoch 1 (head only)", "Best 93.86% at epoch 8 of 12; 48 min on a Colab T4", "v2 figures are platform-reported; local re-score pending"], size=13)


def slide13():
    s = base(13, "Object Detector and Hazard-Exempt Identity Prior", "The detector names the object; the prior down-weights materials that do not fit, never below hazards.")
    picture(s, GAL / "detector_12.png", 0.4, 1.2, 6.9, 5.7)
    bullets(s, 7.5, 1.3, 5.5, 2.6, ["YOLO26n, 2.4 M parameters, TACO litter dataset", "1-class locator: mAP50 0.700;  18-class identifier: 0.383", "4.2 ms per image on the laptop GPU"], size=15, gap=6)
    card(s, 7.5, 3.6, 5.5, 3.2, "Identity prior", ["p'_i = w_i p_i / Σ_j w_j p_j", "w_i = 1 if i ∈ M or i ∈ H, else 0.15", "M: materials that fit the object; H: hazardous classes",
                                                   "Soft weights, not a hard mask; hazards always weight 1; disagreement goes to manual review"], size=14)


def slide14():
    s = base(14, "Uncertainty and the Two-Tier Hazard Gate", "A confident hazard goes to the hazardous bin; an unsure hazard goes to a person. Threshold is set from a review budget.")
    picture(s, REP / "diagram_hazard_gate.png", 0.4, 1.2, 5.5, 5.7)
    card(s, 6.2, 1.3, 6.8, 2.4, "Monte Carlo dropout", ["p̄ = (1/T) Σ_t softmax(f_θ(x; dropout_t)),  T = 25", "U(x) = −Σ_c p̄_c log p̄_c / log C   (0 = sure, 1 = unsure)", "τ_u = Q_{1−β}({U(x_j)}): a share β of items goes to review"], size=14)
    bullets(s, 6.2, 3.9, 6.8, 3.0, ["Test clip of unfamiliar items: false hazardous actions 12 → 0", "Temperature scaling (T = 0.71): ECE 16.6% → 7.3%",
                                    "Honest finding: MC dropout AUROC 0.60 < max-probability 0.66; fixed 0.5 threshold flags 34% of items; calibrate to a review budget instead"], size=15, gap=8)


def slide15():
    s = base(15, "Conformal Family-Set Routing: a Stated Hazard Limit", "Sum probabilities per family, calibrate a threshold per family (stricter for hazards), route from the set of plausible families.")
    picture(s, REP / "diagram_conformal.png", 0.4, 1.2, 4.7, 5.8)
    card(s, 5.3, 1.3, 7.7, 3.0, "Method", ["m_f(x) = Σ_{i∈f} p_i(x)       C(x) = { f : 1 − m_f(x) ≤ q_f }", "q_f = s_(k_f),  k_f = ⌈(n_f + 1)(1 − α_f)⌉   (Beta-corrected rank for a per-deployment guarantee)",
                                              "Hazard in the set → never a recycling bin; automatic action needs a single family and low uncertainty", "Pr[hazard reaches a non-hazard gate] ≤ αH"], size=14)
    rows = [["αH", "Conformal leakage", "Review W", "Matched confidence gate"], ["0.05", "2.84%", "5.7%", "3.50%"], ["0.02", "1.57%", "7.8%", "3.55% (2.3× higher)"], ["0.01", "0.78%", "31.1%", "1.77%"]]
    table(s, 5.3, 4.5, 7.7, rows, [1, 2, 1.5, 2.6], size=13, row_h=0.45)
    tb(s, 5.3, 6.4, 7.7, 0.6, "Top-1 routing with no gate: 6.5% of hazards (48 of 738) reach recycling. Valid for items like the calibration set.", size=13, italic=True)


def slide16():
    s = base(16, "Organic Contamination Index: Two Sensors, Three Models", "A camera cannot see grease; moisture and gas can. A dedicated model is chosen by which sensor works, so sensitivity holds without re-tuning.")
    picture(s, REP / "diagram_oci.png", 0.4, 1.25, 12.5, 2.0)
    card(s, 0.5, 3.45, 6.2, 3.4, "Equations", ["f_m = clip((m_raw − m_dry)/(m_wet − m_dry), 0, 1)", "f_g = clip((−ln(Rs/R0) − l_min)/(l_max − l_min), 0, 1)", "OCI = σ(β0 + β1 f_m + β2 f_g + β3 f_m f_g)",
                                              "τ* = argmin FPR(τ) subject to TPR(τ) ≥ 0.95"], size=14)
    rows = [["Gas channel lost (simulated)", "Sensitivity"], ["Dedicated model, own threshold", "94.9%"], ["Default value 0, shared threshold", "53.7%"], ["Moisture lost: dedicated / default 0", "94.5% / 14.9%"]]
    table(s, 7.0, 3.55, 6.0, rows, [3.8, 1.8], size=13, row_h=0.5)
    tb(s, 7.0, 5.7, 6.0, 1.1, "Simulated sensors only: weakly separable (AUC 0.81 combined, 0.76 single channel). Real calibration is the next stage.", size=13, italic=True)


def slide17():
    s = base(17, "Decision Engine, Explainability, Federated Learning and Video", "Four more components work on the same models.")
    picture(s, REP / "diag_decision_flow.png", 0.3, 1.2, 4.9, 5.8)
    card(s, 5.4, 1.3, 3.7, 2.7, "Explainability", ["Grad-CAM (ConvNeXt stream)", "LIME superpixels, model-agnostic", "SHAP per sensor, exact for linear OCI", "MobileSAM outline removes background"], size=13)
    card(s, 9.3, 1.3, 3.7, 2.7, "Federated learning", ["Units share parameters only", "θ = Σ (n_k/N) θ_k", "5 non-identical units: 89.4% vs 89.0% central vs 81.4% alone"], size=13)
    card(s, 5.4, 4.2, 3.7, 2.6, "Video survey", ["ByteTrack IDs, one entry per item", "240-frame clip: 160 detections → 16 entries", "Hazard-suspected entries go to review"], size=13)
    card(s, 9.3, 4.2, 3.7, 2.6, "Edge export", ["ONNX parity 9.4e-6", "203 MB classifier, 9.3 MB detector", "14 ms GPU, 116 ms CPU on the laptop"], size=13)


def slide18():
    s = base(18, "Results: Classification on 4,232 Held-Out Images", "Item vs family accuracy: the 4.8-point gap is look-alike items that share a bin.")
    rows = [["Metric", "v1", "95% CI", "v2"], ["Item accuracy (33 classes)", "89.30%", "88.42–90.19", "93.83%"], ["Family (routing) accuracy", "94.14%", "93.43–94.80", "97.28%"], ["Hazard recall", "93.50%", "91.63–95.13", "98.37%"],
            ["Macro F1", "0.863", "0.850–0.874", ""], ["Top-3 accuracy", "96.74%", "", ""], ["Macro ROC-AUC", "0.981", "", ""], ["Cohen's kappa", "0.885", "", ""]]
    table(s, 0.5, 1.35, 6.3, rows, [2.6, 1.0, 1.4, 1.0], size=13, row_h=0.46)
    picture(s, ML / "confusion_matrix_families.png", 7.0, 1.25, 6.0, 5.0)
    bullets(s, 0.5, 5.2, 6.3, 1.8, ["7 of the 8 most frequent confusions stay inside one family (cardboard, steel vs aluminium cans, e-waste vs medical)", "Weakest classes are look-alikes (F1 0.56–0.62), not rare classes"], size=14)


def slide19():
    s = base(19, "Analysis: Fit, Calibration and Class Imbalance", "Answers to the usual expert questions: over/underfitting, calibration, imbalance.")
    picture(s, ML / "reliability_diagram.png", 0.3, 1.2, 6.4, 3.0)
    picture(s, ML / "fit_diagnostics.png", 6.7, 1.2, 6.4, 3.0)
    card(s, 0.4, 4.4, 4.1, 2.5, "Over / underfitting", ["Train 96.89%, val 89.51%, test 89.30%", "Moderate gap, not harmful: best epoch = last", "Val ≈ test (0.21 pt)"], size=15)
    card(s, 4.65, 4.4, 4.1, 2.5, "Calibration", ["Under-confident: 73% sure at 89% accuracy", "Temperature 0.71 → ECE 16.6% → 7.3%", "No prediction changes"], size=15)
    card(s, 8.9, 4.4, 4.1, 2.5, "Class imbalance (11.6×)", ["Corrected twice (sampler + weights)", "One correction or logit adjustment: under 1 point", "Hazard recall 93.5% → 94.7% post-hoc"], size=15)


def slide20():
    s = base(20, "Classical Baselines: PCA, LDA, SVM, kNN, Random Forest", "Why deep learning: 71% vs 89-93%. Key finding: SVM on frozen ImageNet features beat the first fine-tuned model.")
    picture(s, REP / "fig_baselines.png", 0.3, 1.2, 8.0, 3.8)
    picture(s, ML / "pca_explained_variance.png", 8.4, 1.2, 4.6, 3.8)
    bullets(s, 0.5, 5.2, 12.4, 1.8, ["Best hand-crafted pipeline (colour + HOG → PCA → RBF SVM) 71.1%; deep features 89–93% (McNemar p ≈ 5e-135)", "Frozen ImageNet + RBF SVM: 93.3% / 97.1% / 98.1% hazard recall, better than the first hybrid (p ≈ 8e-21) → fine-tuning distortion → improved recipe",
                                      "Fine-tuned embedding: 95% of variance in 29 dimensions (715 hand-crafted, 513 frozen)"], size=14, gap=6)


def slide21():
    s = base(21, "Safety Results: Conformal Sets, Gates and Ablation", "Review workload is the price of safety; conformal sets plus a gate give the lowest leakage.")
    picture(s, SIM / "fig_conformal_curve.png", 0.3, 1.2, 6.5, 3.5)
    rows = [["Configuration", "Review W", "Hazard leak"], ["top-1 only", "0.0%", "6.60%"], ["max-probability gate", "10.1%", "3.30%"], ["MC dropout gate", "10.1%", "4.88%"], ["conformal sets", "5.6%", "2.88%"],
            ["conformal + MC gate", "13.0%", "2.64%"], ["conformal + max-prob gate", "12.4%", "2.27%"]]
    table(s, 7.0, 1.3, 6.0, rows, [3.0, 1.4, 1.5], size=13, row_h=0.46)
    bullets(s, 0.5, 4.9, 12.4, 2.1, ["Measured leakage stays below αH for every αH ≤ 0.05 (300 calibration splits)", "Beta-corrected rank: misses the target in 2.1% of calibrations vs 39.5% for the standard rank (exact simulation)",
                                      "Cost: share placed in a correct bin automatically falls from 94.1% (top-1) to 84.6% because items go to a person"], size=14, gap=6)


def slide22():
    s = base(22, "Robustness to Damaged Images and Sensor-Loss Study", "The conformal limit is for items like the calibration set; the gate adds protection under heavy damage. Sensor study is simulated.")
    picture(s, SIM / "fig_distribution_shift.png", 0.3, 1.2, 7.6, 4.0)
    picture(s, SIM / "fig_oci_dropout_synthetic.png", 8.0, 1.2, 5.0, 4.0)
    bullets(s, 0.5, 5.35, 12.4, 1.6, ["Noise σ = 0.2 (item accuracy 53.9%): hazard leakage 28.5% top-1, 15.6% sets alone, 9.8% with the maximum-probability gate",
                                      "Moderate damage (accuracy ≥ 77%): combined mechanism stays at or near the 5% target",
                                      "Sensor loss: 8.6 of 100 decisions flip with the dedicated model vs 48.5 with a default value"], size=14, gap=6)


def slide23():
    s = base(23, "Demonstration 1: Routing Decisions and the Hazardous Family", "Random held-out images; calibration on a separate half. Hazards with a wrong top class still stay out of recycling. All 33 classes are in report Figs. 6.1 to 6.3.")
    picture(s, ROOT / "reports/demo/routing_demo.png", 0.3, 1.2, 8.1, 5.75)
    picture(s, SM / "hazards.jpg", 8.5, 1.2, 4.6, 5.75)


def slide24():
    s = base(24, "Demonstration 2: Videos, Segmentation and Explanations", "Play the two videos: the 33-class tour with routing banner, and the litter survey (160 detections merged into 16 entries).")
    MEDIA = ROOT / "docs/report/media"
    s.shapes.add_movie(str(MEDIA / "class_tour_h264.mp4"), Inches(0.4), Inches(1.25), Inches(4.2), Inches(3.15), poster_frame_image=str(MEDIA / "poster_tour.png"), mime_type="video/mp4")
    s.shapes.add_movie(str(MEDIA / "video_survey_h264.mp4"), Inches(4.75), Inches(1.25), Inches(4.2), Inches(3.15), poster_frame_image=str(MEDIA / "poster_survey.png"), mime_type="video/mp4")
    tb(s, 0.4, 4.42, 4.2, 0.35, "33-class tour video (click to play)", size=12, italic=True, align=PP_ALIGN.CENTER)
    tb(s, 4.75, 4.42, 4.2, 0.35, "Video litter survey (click to play)", size=12, italic=True, align=PP_ALIGN.CENTER)
    picture(s, SM / "sam.jpg", 9.1, 1.25, 4.0, 1.5)
    tb(s, 9.1, 2.75, 4.0, 0.35, "Segment Anything outline and cut-out", size=12, italic=True, align=PP_ALIGN.CENTER)
    picture(s, SM / "gradcam.jpg", 9.1, 3.1, 4.0, 3.9)
    picture(s, SM / "explain_panel.jpg", 0.4, 4.9, 8.55, 2.1)


def slide25():
    s = base(25, "Cost, Schedule, Limitations, Conclusion and Future Work", "Close with the honest status: TRL 3, software validated on real data, hardware next. Thank the guide and invite questions.")
    card(s, 0.4, 1.3, 4.1, 3.0, "Cost and schedule", ["Software: open source, free datasets, ~1 GPU-hour per training run", "Prototype BOM ≈ ₹21,000 (estimate); sensors < ₹1,000", "Timeline: 6 Jul → final review 21 Oct 2026; hardware in Project II"], size=15)
    card(s, 4.6, 1.3, 4.1, 3.0, "Limitations (stated openly)", ["No physical hardware: sensors simulated, Pi latency unmeasured", "Studio photos → field footage differs; conformal limit weakens under heavy damage", "v2 figures platform-reported, local re-score pending"], tcolor=RED, size=15)
    card(s, 8.8, 1.3, 4.2, 3.0, "Future work", ["Build the prototype; real sensor calibration", "Site images + recalibration; ConvNeXt-only test", "Quantisation; federated round on real units"], tcolor=GREEN, size=15)
    card(s, 0.4, 4.5, 12.6, 1.6, "Conclusion", ["Safety and calibration around a strong classifier: 33 classes with 89.3–93.8% accuracy, hazard leakage 6.5% → 1.57% with a stated limit, contamination scoring that survives sensor loss, "
                                                "explanations, private learning and video inventory — validated in software (TRL 3)."], size=16)
    tb(s, 0.4, 6.3, 12.6, 0.6, "Thank you — questions are welcome.", size=24, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


for fn in (slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10, slide11, slide12, slide13, slide14, slide15, slide16, slide17, slide18, slide19, slide20,
           slide21, slide22, slide23, slide24, slide25):
    fn()
assert len(prs.slides) == TOTAL, len(prs.slides)
prs.save(str(OUT))
print("saved", OUT, len(prs.slides), "slides")
