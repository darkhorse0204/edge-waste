# make_report.py - puts the front matter, nine chapters, references and appendix together into the report word file
"""Run:  python scripts/report/make_report.py
Writes docs/report/BITE497J_Project_I_Report_raw.docx (contents fields still empty; run finalize_report.ps1 next)."""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_report as br  # noqa: E402
import report_front as FR  # noqa: E402
import report_refs as RF  # noqa: E402
import report_ch1, report_ch2, report_ch3, report_ch4, report_ch5, report_ch6, report_ch7, report_ch8, report_appendix  # noqa: E402

ROOT = br.ROOT
REP = ROOT / "docs/report/figures"
ML = ROOT / "reports/ml_analysis/figures"
SIM = ROOT / "reports/simulations/figures"
GAL = ROOT / "reports/demo/gallery"
DET = ROOT / "runs/detect/runs/detect"
ASSETS = {
    # diagrams
    "sys": REP / "diagram_system.png", "gate": REP / "diagram_hazard_gate.png", "oci_d": REP / "diagram_oci.png", "conf_d": REP / "diagram_conformal.png",
    "data_pipe": REP / "diag_data_pipeline.png", "hybrid": REP / "diag_hybrid_model.png", "seq": REP / "diag_sequence.png", "hw": REP / "diag_hardware.png",
    "fed_d": REP / "diag_federated.png", "dec": REP / "diag_decision_flow.png", "cls_d": REP / "diag_classes.png", "uc": REP / "diag_usecase.png",
    # new plots
    "lookalike": REP / "fig_lookalike.png", "taxonomy": REP / "fig_taxonomy.png", "counts": REP / "fig_dataset_counts.png", "temp": REP / "fig_temperature.png",
    "gantt": REP / "fig_gantt.png", "baselines": REP / "fig_baselines.png", "versions": REP / "fig_versions.png", "oci_roc": REP / "fig_oci_roc.png", "cost": REP / "fig_cost.png",
    # analysis
    "imbalance": ML / "class_imbalance.png", "cm_fam": ML / "confusion_matrix_families.png", "cm_items": ML / "confusion_matrix_items.png", "embed": ML / "embedding_projections.png",
    "fit": ML / "fit_diagnostics.png", "attn": ML / "fusion_attention.png", "pca": ML / "pca_explained_variance.png", "pcpr": ML / "per_class_precision_recall.png", "rel": ML / "reliability_diagram.png",
    # simulations
    "conf_curve": SIM / "fig_conformal_curve.png", "shift": SIM / "fig_distribution_shift.png", "oci_drop": SIM / "fig_oci_dropout_synthetic.png", "ablation": SIM / "fig_routing_ablation.png",
    "v2hist": SIM / "fig_training_history_v2.png", "unc": SIM / "fig_uncertainty.png",
    # demonstrations
    "routing_demo": ROOT / "reports/demo/routing_demo.png", "explain_panel": ROOT / "reports/demo/explain_panel.png", "video_frames": ROOT / "reports/demo/video/video_frames.png",
    "classes_1": GAL / "classes_1.png", "classes_2": GAL / "classes_2.png", "classes_3": GAL / "classes_3.png", "per_class_f1": GAL / "per_class_f1.png", "hazards": GAL / "hazards.png",
    "gradcam": GAL / "gradcam_16.png", "detector": GAL / "detector_12.png", "sam": GAL / "sam_6.png", "tour_strip": GAL / "class_tour_filmstrip.png", "survey12": GAL / "survey_12_scenes.png",
    "survey_inv": GAL / "survey_inventory.png", "oci_levels": GAL / "oci_levels.png", "federated": GAL / "federated.png", "gan": GAL / "gan.png",
    # detector training outputs
    "det_pr": DET / "taco_multiclass/BoxPR_curve.png", "det_cm": DET / "taco_multiclass/confusion_matrix_normalized.png", "det_results": DET / "taco_multiclass/results.png",
}
JPEG = {"classes_1", "classes_2", "classes_3", "hazards", "gradcam", "detector", "survey12", "routing_demo", "video_frames", "tour_strip", "explain_panel", "sam"}
CHAPTERS = [report_ch1, report_ch2, report_ch3, report_ch4, report_ch5, report_ch6, report_ch7, report_ch8]


def slim_assets():
    """Re-save the large screen captures as optimised jpeg copies so the word file stays a reasonable size."""
    from PIL import Image
    out = ROOT / "docs/report/figures_small"
    out.mkdir(exist_ok=True)
    for k in JPEG:
        src = ASSETS[k]
        im = Image.open(src).convert("RGB")
        if im.width > 1800:
            im = im.resize((1800, int(im.height * 1800 / im.width)), Image.LANCZOS)
        dst = out / (k + ".jpg")
        im.save(dst, quality=85, optimize=True)
        ASSETS[k] = dst


def centered(rep, text, size, bold=False, italic=False, before=0, after=6, caps=False):
    p = rep.d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.0
    r = p.add_run(text.upper() if caps else text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    return p


def cover(rep):
    centered(rep, FR.TITLE, 20, bold=True, before=30, after=12)
    centered(rep, "BITE497J – Project I", 18, after=40)
    centered(rep, "Submitted in partial fulfillment of the requirements for the degree of", 14, italic=True, after=36)
    centered(rep, "Bachelor of Technology", 20, bold=True, after=8)
    centered(rep, "in", 14, after=8)
    centered(rep, "Information Technology", 18, bold=True, after=36)
    centered(rep, "by", 14, italic=True, after=10)
    for name, reg in FR.STUDENTS:
        centered(rep, f"{name} – {reg}", 16, bold=True, caps=True, after=4)
    centered(rep, "Under the guidance of", 16, before=26, after=4)
    centered(rep, FR.GUIDE, 16, bold=True, after=34)
    centered(rep, "School of Computer Science Engineering and Information Systems", 12, bold=True, after=2)
    centered(rep, "VIT, Vellore", 12, bold=True, after=18)
    p = rep.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(ROOT / "docs/report/assets/vit_logo.png"), width=Cm(6.2))
    centered(rep, FR.DATE, 12, bold=True, after=0)


def body_par(rep, text, first=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = rep.d.add_paragraph()
    p.alignment = align
    if not first:
        p.paragraph_format.first_line_indent = Cm(1.27)
    rep.runs(p, text)
    return p


def left_par(rep, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, before=0, line=1.15):
    p = rep.d.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before); p.paragraph_format.line_spacing = line
    rep.runs(p, text, bold=bold)
    return p


def front_heading(rep, text, toc=True, page_break=True):
    p = rep.d.add_paragraph(style="Front Heading" if toc else "Front Heading Plain")
    p.paragraph_format.page_break_before = page_break
    p.add_run(text)
    return p


def names_line():
    return "; ".join(f"{n} ({r})" for n, r in FR.STUDENTS)


def main():
    slim_assets()
    rep = br.Report()
    numbered = [(str(i), mod) for i, mod in enumerate(CHAPTERS, start=1)]
    all_blocks = {str(i): m.blocks() for i, m in numbered}
    all_blocks["A"] = report_appendix.blocks()
    rep.number_blocks([(k, {"blocks": v}) for k, v in all_blocks.items()])
    d = rep.d

    # ------------------------------------------------------------ section 1: cover (number hidden)
    s0 = d.sections[0]
    rep.page_setup(s0)
    rep.set_pgnum(s0, "lowerRoman", 1)
    rep.footer_page_number(s0, visible=False)
    cover(rep)

    # ------------------------------------------------------------ section 2: front matter (roman numbers)
    s1 = d.add_section(WD_SECTION.NEW_PAGE)
    rep.page_setup(s1)
    rep.set_pgnum(s1, "lowerRoman")
    rep.footer_page_number(s1, visible=True)
    cover(rep)

    front_heading(rep, "Declaration", toc=False)
    left_par(rep, "", after=6)
    body_par(rep, f"We hereby declare that the BITE497J – Project I thesis entitled “{FR.TITLE}” submitted by us, for the award of the degree of <b>Bachelor of Technology in "
                  f"Information Technology</b>, School of Computer Science Engineering and Information Systems to VIT is a record of bonafide work carried out by us under the supervision of "
                  f"<b>{FR.GUIDE_FULL}, SCORE, VIT, Vellore</b>.")
    body_par(rep, "We further declare that the work reported in this thesis has not been submitted and will not be submitted, either in part or in full, for the award of any other degree "
                  "or diploma in this institute or any other institute or university.")
    left_par(rep, "", after=30)
    left_par(rep, "Place: Vellore")
    left_par(rep, "Date:")
    left_par(rep, "", after=24)
    for n, r in FR.STUDENTS:
        left_par(rep, f"{n} ({r})", align=WD_ALIGN_PARAGRAPH.RIGHT)
    left_par(rep, "Signature of the Candidates", align=WD_ALIGN_PARAGRAPH.RIGHT, before=6)

    front_heading(rep, "Certificate", toc=False)
    left_par(rep, "", after=6)
    body_par(rep, f"This is to certify that the BITE497J – Project I thesis entitled “{FR.TITLE}” submitted by {names_line()}, SCORE, VIT, for the award of the degree of "
                  "<b>Bachelor of Technology in Information Technology</b>, School of Computer Science Engineering and Information Systems, is a record of bonafide work carried out by "
                  "them under my supervision during the period 06.07.2026 to 03.11.2026, as per the VIT code of academic and research ethics.")
    body_par(rep, "The contents of this report have not been submitted and will not be submitted either in part or in full, for the award of any other degree or diploma in this institute "
                  "or any other institute or university. The thesis fulfills the requirements and regulations of the University and in my opinion meets the necessary standards for submission.")
    left_par(rep, "", after=24)
    left_par(rep, "Place: Vellore")
    left_par(rep, "Date:")
    left_par(rep, "", after=20)
    left_par(rep, "Signature of the Guide", align=WD_ALIGN_PARAGRAPH.RIGHT)
    left_par(rep, "", after=30)
    p = left_par(rep, "Internal Examiner")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(br.TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\tExternal Examiner")
    left_par(rep, "", after=40)
    left_par(rep, "Head of the Department", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    left_par(rep, "Department of Information Technology", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    front_heading(rep, "Acknowledgement")
    left_par(rep, "", after=6)
    body_par(rep, f"It is our pleasure to express with a deep sense of gratitude to our BITE497J – Project I guide <b>{FR.GUIDE_FULL}</b>, School of Computer Science Engineering and "
                  "Information Systems, Vellore Institute of Technology, Vellore for constant guidance, continual encouragement and support in our endeavor. Our association with the guide "
                  "is not confined to academics only, but it is a great opportunity on our part to work with an intellectual and an expert in the field of "
                  f"<b>{FR.DOMAIN}</b>.")
    body_par(rep, "We would like to express our heartfelt gratitude to Honorable Chancellor <b>Dr. G Viswanathan</b>; respected Vice Presidents <b>Dr. Sankar Viswanathan</b>, "
                  "<b>Dr. Sekar Viswanathan</b>, Vice Chancellor <b>Dr. V. S. Kanchana Bhaaskaran</b>; Pro-Vice Chancellor <b>Dr. Partha Sharathi Mallick</b>; and Registrar "
                  "<b>Dr. Jayabarathi T.</b>")
    body_par(rep, "Our whole-hearted thanks to Dean <b>Dr. Daphne Lopez</b>, School of Computer Science Engineering and Information Systems, Head, Department of Information Technology, "
                  "<b>Dr. Arivuselvan K</b>, Information Technology Project Coordinator <b>Dr. Suganya P</b>, SCORE School Project Coordinator <b>Dr. Karthikeyan P</b>, all faculty, staff and "
                  "members working as limbs of our university for their continuous guidance throughout our course of study in unlimited ways.")
    body_par(rep, "It is indeed a pleasure to thank our parents and friends who persuaded and encouraged us to take up and complete our project successfully. Last, but not least, we express "
                  "our gratitude and appreciation to all those who have helped us directly or indirectly towards the successful completion of the project.")
    left_par(rep, "", after=12)
    left_par(rep, "Place: Vellore")
    left_par(rep, "Date:")
    for n, r in FR.STUDENTS:
        left_par(rep, n, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True)

    front_heading(rep, "Executive Summary")
    left_par(rep, "", after=6)
    assert len(FR.EXEC_SUMMARY.split()) <= 200
    body_par(rep, FR.EXEC_SUMMARY, first=True)

    front_heading(rep, "Table of Contents")
    p = left_par(rep, "Chapter and section", bold=True, after=4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(br.TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\tPage No.").bold = True
    rep.field_paragraph('TOC \\o "1-2" \\u \\h \\z', "Right-click here and choose Update Field to build the table of contents.")

    front_heading(rep, "List of Figures")
    p = left_par(rep, "Figure number and title", bold=True, after=4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(br.TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\tPage No.").bold = True
    rep.field_paragraph('TOC \\h \\z \\t "Figure Caption,3"', "Right-click here and choose Update Field to build the list of figures.")

    front_heading(rep, "List of Tables")
    p = left_par(rep, "Table number and title", bold=True, after=4)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(br.TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\tPage No.").bold = True
    rep.field_paragraph('TOC \\h \\z \\t "Table Caption,4"', "Right-click here and choose Update Field to build the list of tables.")

    front_heading(rep, "List of Abbreviations")
    left_par(rep, "", after=4)
    for ab, full in FR.ABBREVIATIONS:
        p = left_par(rep, "", after=2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.4))
        p.paragraph_format.left_indent = Cm(3.4); p.paragraph_format.first_line_indent = Cm(-3.4)
        p.add_run(ab).bold = True
        p.add_run("\t" + full)

    front_heading(rep, "Symbols and Notations")
    left_par(rep, "", after=4)
    for sym, meaning in FR.SYMBOLS:
        p = left_par(rep, "", after=2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.4))
        p.paragraph_format.left_indent = Cm(3.4); p.paragraph_format.first_line_indent = Cm(-3.4)
        rep.runs(p, sym, italic=True)
        p.add_run("\t")
        rep.runs(p, meaning)

    # ------------------------------------------------------------ section 3: chapters 1-9 (arabic numbers from 1)
    s2 = d.add_section(WD_SECTION.NEW_PAGE)
    rep.page_setup(s2)
    rep.set_pgnum(s2, "decimal", 1)
    rep.footer_page_number(s2, visible=True)
    for i, mod in numbered:
        rep.heading(mod.TITLE, 1)
        rep.render_blocks(all_blocks[i], ASSETS)
    # chapter 9: references
    rep.heading(("CHAPTER 9", "REFERENCES"), 1)
    pp = d.add_paragraph(style="Normal"); pp.paragraph_format.first_line_indent = Cm(1.27)
    pp.add_run("The references follow the APA 7th edition. Citations in the text give the author and the year. The list is in alphabetical order, followed by patents, websites and data sets.")
    for r in RF.sorted_refs(RF.REFS):
        p = d.add_paragraph(style="Reference"); rep.runs(p, r)
    for title, lst in (("Patents", RF.PATENTS), ("Websites", RF.WEBSITES), ("Data Sets", RF.DATASETS)):
        h = d.add_paragraph(style="Heading 2"); h.add_run(title)
        for r in RF.sorted_refs(lst):
            p = d.add_paragraph(style="Reference"); rep.runs(p, r)

    # ------------------------------------------------------------ section 4: appendix (no page numbers)
    s3 = d.add_section(WD_SECTION.NEW_PAGE)
    rep.page_setup(s3)
    rep.set_pgnum(s3, "decimal")
    rep.footer_page_number(s3, visible=False)
    rep.heading(report_appendix.TITLE, 1)
    rep.render_blocks([("h3", b[1]) if b[0] == "h2" else b for b in all_blocks["A"]], ASSETS)

    br.OUT.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(br.OUT))
    print("saved", br.OUT)


if __name__ == "__main__":
    main()
