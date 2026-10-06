# build_dossier.py - writes the technical dossier (simple english) for the project review into review/
"""usage: python scripts/report/build_dossier.py   ->  review/Technical_Dossier.docx"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "review" / "Technical_Dossier.docx"
FIG = ROOT / "docs" / "report" / "figures"
ML = ROOT / "reports" / "ml_analysis" / "figures"
SIM = ROOT / "reports" / "simulations" / "figures"

d = Document()
s = d.sections[0]
s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2.0)
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)
d.styles["Normal"].paragraph_format.space_after = Pt(5)
for n, sz in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 12)):
    h = d.styles[n]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    rf = h.element.get_or_add_rPr().find(qn("w:rFonts"))
    for a in ("w:asciiTheme", "w:hAnsiTheme"):
        rf.attrib.pop(qn(a), None)
    rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")


def shade(cell, fill):
    p = cell._tc.get_or_add_tcPr(); e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), fill); p.append(e)


def H(t, l=1):
    d.add_heading(t, l)


def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    return p


def B(items):
    for i in items:
        d.add_paragraph(i, style="List Bullet")


def TBL(header, rows, widths, size=9.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); shade(c, "D9E2F3")
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ""; cs[i].paragraphs[0].add_run(str(v)).font.size = Pt(size)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    d.add_paragraph()


def FIGURE(path, cap, w=15.0):
    d.add_picture(str(path), width=Cm(w))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True; c.runs[0].font.size = Pt(9.5)


def CODEBOX(text, size=8.5):
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.rows[0].cells[0]; shade(c, "F2F2F2"); c.text = ""
    first = True
    for ln in text.split("\n"):
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln or " "); r.font.name = "Consolas"; r.font.size = Pt(size)
    d.add_paragraph()


# ------------------------------------------------------------------------------------------------ title
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Technical Dossier"); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.add_run("Edge-Deployed Waste Classification with Sensor-Augmented Vision (CAPS-FCL)").font.size = Pt(15)
t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.add_run("BITE497J Project I  |  Ansh Jerath (23BIT0123), Arnav Mishra (23BIT0142), Bhavesh Singh Thakur (23BIT0199)  |  Guide: Dr. Valarmathi B.").font.size = Pt(10.5)
P("Written in simple English for the review. Every number in this dossier comes from a saved result in the project folder. "
  "Where something is not done or not proven, it says so.", italic=True)

H("How to use this dossier")
B(["Section 1 to 3: the idea in plain words. Read these first.",
   "Section 4 to 9: data, model, training, results, analysis and the safety parts, each with what, how, why and where the code is.",
   "Section 10 to 12: extra parts, the hardware prototype, and the honest limits.",
   "Section 13: a map from every feature to the file that codes it. Section 14: how to run the demo in 10 minutes. Section 15: questions the professor may ask, with answers. Section 16: words explained."])

# ------------------------------------------------------------------------------------------------ 1
H("1. What the project does (in plain words)")
P("A waste sorting machine looks at one item, decides what it is, and sends it to the right bin. Most machines use one camera and one model, give one answer, and never say how sure they are. "
  "This project builds a smarter sorter in software:")
B(["It recognises 33 kinds of items, grouped into 9 material families (plastic, paper, cardboard, glass, metal, organic, styrofoam, textile, hazardous). One family = one bin.",
   "It says how sure it is. If it is unsure, a person checks the item.",
   "It never sends a dangerous item (battery, electronic waste, medical waste) to recycling. The operator can set how rare a mistake must be, and the system keeps to it for items like the ones it was calibrated on.",
   "Two cheap sensors (humidity or moisture, and gas) tell if an item is dirty with food, which a camera cannot see.",
   "It can explain its decisions with heat maps, and it can learn together with other machines without sharing photos (federated learning)."])
P("Status in one line: the software is built and tested on real held-out images. The sensors are simulated, and a small physical prototype with an ESP32 is planned. This is Technology Readiness Level 3 (proof of concept).", bold=True)

H("2. The problem and the goals")
TBL(["Problem in existing sorters", "What we did"], [
    ("one model, one answer, no confidence", "uncertainty with monte carlo dropout and calibrated probabilities"),
    ("only a few broad classes", "33 classes, 9 families, and results reported at item level and family level"),
    ("a wrong hazard call is treated like any other mistake", "hazard-safe rules: two-tier gate and conformal family sets with a stated limit"),
    ("camera cannot see food residue", "organic contamination index from two sensors, with a model for each sensor case"),
    ("sending photos to a server costs bandwidth and privacy", "federated averaging: only model weights move (simulated)"),
], [8.0, 8.4])

# ------------------------------------------------------------------------------------------------ 3
H("3. The big picture")
FIGURE(FIG / "diagram_system.png", "Figure 1. System: camera, two models, cross-check, uncertainty, family sets, sensors, decision engine", 13.5)
P("What happens to one item, step by step:", bold=True)
B(["1. The camera takes a picture. The detector (YOLO) finds the item and names the object (for example: can).",
   "2. The classifier (ConvNeXt + Vision Transformer) gives a probability for each of the 33 classes.",
   "3. The cross-check lowers materials that do not fit the object. Hazardous classes are never lowered.",
   "4. Uncertainty is measured (monte carlo dropout). A set of possible families is built (conformal prediction).",
   "5. The sensors give a contamination score.",
   "6. The decision engine picks one action: hazardous bin, priority review, contamination reject, or the bin of the family."])
FIGURE(FIG / "diag_decision_flow.png", "Figure 2. Decision flow", 10.0)

# ------------------------------------------------------------------------------------------------ 4
H("4. Data")
P("We use three public Kaggle datasets (the notebook downloads them). The 30-class household dataset is the main part. TrashBox is used only for electronic waste and medical waste, and Garbage Classification 12 gives battery, clothing, shoes and extra food waste. "
  "Folders that are too broad (for example one big 'plastic' folder) are dropped, because they cannot be put in one of our fine classes without guessing.")
FIGURE(FIG / "fig_taxonomy.png", "Figure 3. The 33 classes in 9 families, with image counts", 15.0)
TBL(["Item", "Value"], [("total images", "28,203"), ("train / validation / test", "19,739 / 4,232 / 4,232 (70 / 15 / 15, stratified, seed 42)"),
                          ("class imbalance", "largest class is 11.6 times the smallest (clothing 5,825 vs about 500 for most classes)"),
                          ("code", "taxonomy.py (classes and sources), ingest_raw_datasets.py, make_data_splits.py")], [5.0, 11.4])
P("Why stratified: every class keeps the same share in train, validation and test, so rare classes are tested too.")
P("Augmentation (training images only): resize to 256, random crop 224, flip, rotation up to 20 degrees, colour change, blur, random erasing. This copies camera pose, light and partly hidden items. Validation and test are never changed. Code: build_transforms in data/waste_image_dataset.py.")
H("Two data bugs we found and fixed", 2)
B(["Split leak: the file name hash used the operating system path, so Windows and Colab made different splits. Scoring the Colab model on a split built on Windows gave 94.52%, which was too high because about 70% of the 'test' images had been used for training. We found it because the number went up. Fixed with POSIX paths, and the exact Colab split was recovered (the model scores exactly 3,779 of 4,232 = 89.30% on it).",
   "Stale duplicates: one Colab run showed 1,234 steps per epoch instead of 617, because old copies of the images were still on the drive. The split now ignores files not listed in the latest provenance file, and the notebook has a gate that stops if sizes differ."])

# ------------------------------------------------------------------------------------------------ 5
H("5. The model")
FIGURE(FIG / "diag_hybrid_model.png", "Figure 4. Hybrid classifier", 7.5)
P("Two pretrained networks look at the same 224 by 224 picture:")
B(["ConvNeXt-Tiny (27.8 million parameters): a convolutional network, good at texture (foil, cardboard fibres).",
   "Vision Transformer Small (21.7 million parameters): looks at the whole picture with attention, good at shape and context.",
   "Each gives a feature vector, projected to 512 numbers (linear layer, layer norm, GELU).",
   "An attention layer gives each stream a weight (they add up to 1), so the model can trust one more for a given image. We measured it: ConvNeXt gets 82% to 93% of the weight for every family.",
   "The two weighted vectors are joined (1,024 numbers) and a small head (dropout 0.2, 1024 to 512, GELU, dropout 0.2, 512 to 33) gives the 33 scores. Total 50.75 million parameters."])
P("The attention in words: score = a small network applied to each stream; weight = softmax of the two scores (so weights are positive and add to one).", italic=True)
P("Why these choices:", bold=True)
B(["Transfer learning: training from scratch gave only 49% in an earlier try (on an earlier dataset), so we start from ImageNet weights.",
   "GELU: smooth activation used by both pretrained networks; no dead neurons.",
   "LayerNorm instead of BatchNorm: the result does not depend on batch size, so it behaves the same on the edge with one image.",
   "Dropout 0.2 in the head: regularisation, and it is also what monte carlo dropout uses later.",
   "Code: classification/hybrid_model.py, classes AttentionFusion and HybridConvNeXtViT."])

# ------------------------------------------------------------------------------------------------ 6
H("6. Training")
TBL(["Setting", "Value", "Why"], [
    ("optimiser", "AdamW, learning rate 3e-4, weight decay 0.05", "adaptive steps suit pretrained networks plus a new head; weight decay is applied correctly in AdamW"),
    ("learning rate schedule", "one-cycle: warm-up for 1 epoch, then slow decay", "warm-up protects pretrained weights"),
    ("loss", "cross-entropy with label smoothing 0.1", "stops over-confidence between look-alike classes"),
    ("batch size / epochs", "32 / up to 25 with early stopping (patience 5)", "stop when validation stops improving"),
    ("mixed precision, gradient clip 1.0", "on", "speed on the GPU; stability"),
    ("class imbalance", "balanced sampler only (improved recipe)", "the first recipe used sampler and loss weights together, which corrects twice"),
    ("backbones", "frozen for 2 epochs, then 10 times lower learning rate than new layers (3e-5)", "protects pretrained features (see below)"),
], [4.0, 6.2, 6.2])
H("The improved recipe and why it was needed", 2)
P("Our first model reached 89.30%. A simple SVM on frozen ImageNet features reached 93.3%, which is better. That told us fine-tuning had damaged the pretrained features: the new head was random, so large noisy gradients flowed into the good backbones from step one. "
  "The fix: train the head first with frozen backbones (2 epochs), then fine-tune the backbones slowly, and correct imbalance only once. Result on Colab: 93.83% item accuracy.")
H("The real Colab run (executed notebook)", 2)
B(["Tesla T4 GPU, 617 steps per epoch, about 4 minutes per epoch.",
   "Validation accuracy 0.900 after epoch 1 (only the head trained), best 0.939 at epoch 8, training stopped after epoch 13 by early stopping, total about 52 minutes.",
   "Training accuracy ends at 0.975 and validation near 0.938: a small gap. This is mild overfitting, not harmful.",
   "Test set: item accuracy 93.83%, family accuracy 97.28%, hazard recall 98.37% (726 of 738).",
   "The executed notebook is notebooks/colab_classifier_training_executed.ipynb; the annotated copy with comments is review/notebooks/06_colab_training_executed_annotated.ipynb. Training code: classification/train_classifier.py, function train."])
FIGURE(SIM / "fig_training_history_v2.png", "Figure 5. Training history of the improved recipe", 14.0)

# ------------------------------------------------------------------------------------------------ 7
H("7. Results")
TBL(["Metric (4,232 test images)", "First model", "Improved recipe (Colab)"], [
    ("item accuracy (33 classes)", "89.30%", "93.83%"), ("family accuracy (9 bins)", "94.14%", "97.28%"), ("hazard recall", "93.50% (690 of 738)", "98.37% (726 of 738)"),
    ("macro F1", "0.863", "0.9065 (from the Colab log)"), ("top-3 accuracy", "96.74%", "not computed"),
], [6.0, 5.0, 5.4])
P("The detailed analysis (calibration, imbalance, PCA, LDA, SVM, bootstrap intervals, confusion matrices) was made on the first model, because its test split was recovered and verified. "
  "The improved recipe's numbers are from the Colab run itself. A local re-score gave 96.8%, which is too high and is a leak signature (the local split differs from Colab's), so it was discarded. The evaluation command now refuses a split that does not reproduce the validation accuracy stored in the checkpoint.", italic=True)
FIGURE(FIG / "fig_versions.png", "Figure 6. Item accuracy, family accuracy and hazard recall", 11.0)
P("Why two accuracies? A bin is chosen by family. Mixing two kinds of plastic bottle is a recognition error but not a sorting error. Seven of the eight most common mistakes of the first model are inside one family (cardboard boxes and packaging, steel and aluminium food cans, electronic waste and medical waste).")
FIGURE(ML / "confusion_matrix_families.png", "Figure 7. Confusion matrix at family level (first model)", 10.5)
H("Detector", 2)
TBL(["Variant", "Precision", "Recall", "mAP50", "mAP50-95"], [("1-class waste item", "0.871", "0.620", "0.700", "0.492"), ("18-class object", "0.603", "0.361", "0.383", "0.289")], [5.0, 2.8, 2.8, 2.8, 3.0])
P("YOLO26n, 2.4 million parameters, trained on the TACO litter dataset. 4.2 ms per image on the laptop GPU.")

# ------------------------------------------------------------------------------------------------ 8
H("8. Questions about the machine learning (answered with evidence)")
H("Is the model overfitting or underfitting?", 3)
P("First model: train 96.89%, validation 89.51%, test 89.30%. Not underfitting (high training accuracy). A gap of 7.6 points, expected for 50 million parameters and 19.7 thousand images. Not harmful: validation never turned down and validation and test differ by only 0.21 points. "
  "Learning curve: more data of the same kind gives little gain now. Validation curve: the same model underfits at a strong regularisation and overfits at a weak one, the usual bias-variance picture.")
FIGURE(ML / "fit_diagnostics.png", "Figure 8. Fit diagnostics", 14.5)
H("Is the confidence honest? (calibration)", 3)
P("The first model was under-confident: right 89% of the time but only 73% sure on average. One number, the temperature T = 0.71 (fitted on validation), lowers the calibration error (ECE) from 16.6% to 7.3% and changes no prediction. Code: analysis/calibration.py.")
FIGURE(ML / "reliability_diagram.png", "Figure 9. Reliability diagram before and after temperature scaling", 14.0)
H("Is class imbalance handled?", 3)
P("Largest class 11.6 times the smallest. The first run corrected it twice (sampler and loss weights), which over-weights rare classes (about 1 over n squared). We found this and now correct once. "
  "Adjusting the scores after training (score + tau * log class size, tau chosen on validation) moved accuracy from 89.30% to 89.60% and hazard recall from 93.5% to 94.7%: a small but real effect.")
H("Are PCA, LDA, SVM done?", 3)
P("Yes. Seven classical classifiers (naive Bayes, LDA, logistic regression, linear SVM, RBF SVM, kNN, random forest) were trained after PCA on three feature sets (hand-made colour and HOG features, frozen ImageNet features, our fine-tuned features), with the same protocol as the deep model (hyperparameters on validation, test scored once). Code: analysis/classical_baselines.py.")
FIGURE(FIG / "fig_baselines.png", "Figure 10. Test accuracy of the classical classifiers", 14.5)
B(["Hand-made features: best 71.1% (RBF SVM). Deep features: 89% to 93%. So deep learning is needed.",
   "Naive Bayes and LDA: small train-test gap but low accuracy (high bias). Random forest and kNN: about 99% train accuracy but big drop on test (high variance).",
   "RBF SVM on frozen ImageNet features: 93.3% item, 97.1% family, 98.1% hazard recall. This beat our first fine-tuned model and led to the improved recipe.",
   "PCA: the fine-tuned embedding keeps 95% of its information in 29 numbers (715 for hand-made features, 513 for frozen features)."])
P("Bootstrap 95% intervals for the first model: item 88.42 to 90.19%, family 93.43 to 94.80%, hazard recall 91.63 to 95.13%. We did not run k-fold cross validation for the deep model, because one run costs about an hour of GPU; the bootstrap interval on 4,232 test images already shows the uncertainty.")

# ------------------------------------------------------------------------------------------------ 9
H("9. The safety mechanisms (the main ideas)")
H("9.1 Uncertainty (monte carlo dropout)", 2)
P("Idea: run the classifier head 25 times with dropout switched on. If the answers agree, the model is sure. U = entropy of the average answer divided by log(33), so U is between 0 (sure) and 1 (unsure). "
  "The backbones have no dropout, so repeating only the head gives exactly the same result as repeating the whole network, about 100 times cheaper. Code: decision_engine/mc_dropout_uncertainty.py.")
P("Honest result: as a detector of wrong answers, monte carlo dropout (AUROC 0.598) was not better than the plain highest probability (0.657). A fixed threshold of 0.5 sent 34% of items to review. So the threshold is set from a review budget (for example the 10% most unsure calibration items), not fixed.")
H("9.2 Two-tier hazard gate", 2)
P("A confident hazard goes to the hazardous bin. An unsure hazard goes to priority review and never to recycling. On a test clip of unfamiliar items, the unsure classifier called 12 frames hazardous by chance; with the gate, automatic hazardous actions fell from 12 to 0.")
FIGURE(FIG / "diagram_hazard_gate.png", "Figure 11. Two-tier hazard gate", 10.0)
H("9.3 Conformal family-set routing (a stated limit)", 2)
P("Problem: a battery may be read as office paper, and then the single top answer sends it to recycling. Idea: add the class probabilities into family totals, find a threshold for each family on labelled calibration data (stricter for hazardous), and build a SET of possible families for each new item. "
  "If 'hazardous' is in the set, the item cannot go to a recycling bin. Then the chance that a hazard reaches a non-hazard bin is at most alpha_H (the operator's choice) for items that look like the calibration items.")
TBL(["alpha_H", "Hazard leakage (conformal sets)", "Items sent to review", "Matched confidence gate"], [("0.05", "2.84%", "5.7%", "3.50%"), ("0.02", "1.57%", "7.8%", "3.55%"), ("0.01", "0.78%", "31.1%", "1.77%")], [3.0, 5.0, 4.2, 4.2])
P("Without any gate, 6.5% of hazards (48 of 738) reach recycling. At about 8% review, conformal sets leak 2.3 times less than a tuned confidence gate. Measured leakage stays below alpha_H for every alpha_H of 0.05 or less (300 random splits). "
  "Limit: the guarantee does not hold for very different images. With heavy noise (item accuracy 53.9%), leakage rises to 15.6% for sets alone, and an extra uncertainty gate brings it to 9.8% (top-1 alone gives 28.5%). Code: decision_engine/conformal_routing.py.")
FIGURE(SIM / "fig_conformal_curve.png", "Figure 12. Leakage against review workload", 14.0)
H("9.4 Object-material cross-check", 2)
P("The detector says what the object is. A can is not plastic, so plastic gets weight 0.15 and fitting materials weight 1; then the probabilities are renormalised. Hazardous classes always keep weight 1, so a hazard call can never be hidden by this step. "
  "If the two stages still disagree, the item goes to manual review. Example from the test set: an aluminium soda can first read as a plastic bottle (69%); after the cross-check the can rises from 4% to 17% and the bottle falls to 40%. Code: decision_engine/object_identity_prior.py.")
H("9.5 Organic contamination index (OCI)", 2)
P("A greasy pizza box is cardboard but cannot be recycled, and a camera cannot see grease. Moisture and gas readings are scaled to 0..1 with fixed anchors measured once (dry and wet reference; clean air and smelly reference). "
  "Three small logistic models are fitted: both sensors, moisture only, gas only. If a sensor fails, the matching model is used. The threshold is the one with the lowest false alarms among those that find at least 95% of contaminated items (missing dirt costs more than an extra check).")
P("Result on SIMULATED sensors: with the gas channel lost, the dedicated moisture model keeps 94.9% sensitivity; putting zero for the missing gas value in the big model drops it to 53.7%. The simulated sensors are only weakly separable (false positive rate about 80% at 95% sensitivity), so real calibration is essential. Code: contamination/oci_model.py, sensor_normalization.py.")
FIGURE(FIG / "diagram_oci.png", "Figure 13. Organic contamination index", 15.0)
H("9.6 Decision engine", 2)
TBL(["Condition", "Action"], [("hazardous is the only family in the set and U is low", "hazardous bin"), ("hazardous is in the set in any other case", "priority review, never recycling"),
                              ("object and material contradict", "manual review"), ("set empty or several families, or U high", "manual review"),
                              ("index above threshold on a recyclable", "contamination reject"), ("otherwise", "bin of the one family")], [9.5, 6.9])
P("Code: decision_engine/routing_rules.py, function decide. The actuator is a mock (it prints the action) until hardware exists.")

# ------------------------------------------------------------------------------------------------ 10
H("10. Other parts")
B(["Explainability: Grad-CAM (heat map on the ConvNeXt stream), LIME (hide parts of the image and see what changes), SHAP (share of each sensor in the contamination score; adds up exactly). Code: explainability/ and contamination/shap_explanation.py.",
   "Segmentation: MobileSAM draws the exact outline of the item inside the detector box, so background does not reach the classifier. Code: detection/sam_segmentation.py.",
   "Synthetic images: a small GAN (DCGAN) can make extra images for rare classes. It was built and tried on textile; not used in the reported models, and image quality was not measured. Code: augmentation/gan_augmentation.py.",
   "Federated learning: units train locally and send only weights; the server averages them with weights proportional to the number of images. Simulation on saved features, 5 units with different class mixes: shared model 89.4% against 89.0% for one central model and 81.4% for units alone. Code: federated/fedavg_simulation.py (full model) and the review notebook 05 (features only).",
   "Video survey: ByteTrack gives each object an id, so one object is counted once. On a 240-frame clip made from 12 photographs, 160 detections became 16 inventory entries (the clip has no true count, so this shows merging, not counting accuracy). Code: applications/video_litter_survey.py.",
   "Edge export: both models exported to ONNX; outputs match the originals within 9.4e-6. Latency on the laptop: 14 ms on the GPU, 116 ms on the CPU. The Raspberry Pi was not measured."])

# ------------------------------------------------------------------------------------------------ 11
H("11. Hardware prototype (planned, ESP32)")
P("A bench unit: ESP32, DHT22 (humidity, used as the moisture channel), MQ-135 gas sensor, one servo that tilts a tray (left = auto-sorted, right = hazardous, level = ask a person). "
  "The laptop runs the models and sends commands over USB serial. Parts cost roughly 1,000 to 1,800 rupees (estimate). Files: hardware/esp32/sorter_node_esp32/sorter_node_esp32.ino, scripts/prototype_sorter.py, scripts/calibrate_prototype_oci.py, scripts/check_node.py, and the guide docs/ESP32_Prototype_Hardware_Guide.docx. "
  "Status: the software path is tested with a simulated node and real images. The firmware is not compiled or run on hardware yet, and the sensors are not calibrated on real items.")
FIGURE(FIG / "diag_hardware.png", "Figure 14. Planned hardware", 15.0)

# ------------------------------------------------------------------------------------------------ 12
H("12. Honest limits (what is not done)")
B(["No physical sensors or conveyor yet: the contamination index is fitted on simulated data.",
   "Training photos are mostly studio or household photos. On outdoor litter the model is unsure and sends most items to a person. That is safe but shows a domain gap.",
   "The conformal limit is for items like the calibration items; under heavy damage it is not kept, and the uncertainty gate only reduces the risk.",
   "The improved model's detailed analysis (calibration, conformal) was not repeated: it was done on the first model. Its test numbers are from the Colab run.",
   "Monte carlo dropout is not better than the plain highest probability at finding mistakes.",
   "We did not test a ConvNeXt-only model, so we cannot claim that the hybrid is better than ConvNeXt alone. The attention weights only show that the model leans on ConvNeXt.",
   "Federated learning is a single-process simulation. Raspberry Pi latency is not measured.",
   "Swin Transformer was tried for only 2 epochs (83.0% item accuracy), so it is not a fair comparison."])

# ------------------------------------------------------------------------------------------------ 13
H("13. Which feature is coded where")
P("This table is made from the first-line comment of every source file (each file starts with a one-line comment that says what it is).", italic=True)
rows = []
for f in sorted((ROOT / "src" / "edgewaste").rglob("*.py")):
    if f.name == "__init__.py":
        continue
    first = f.read_text(encoding="utf-8").splitlines()[0].lstrip("# ").strip()
    rows.append((str(f.relative_to(ROOT / "src" / "edgewaste")).replace("\\", "/"), first.split(" - ", 1)[-1]))
TBL(["file in src/edgewaste/", "what it does"], rows, [6.2, 10.2], size=8.5)
H("Scripts and notebooks", 2)
rows = []
for f in sorted((ROOT / "scripts").glob("*.py")):
    first = f.read_text(encoding="utf-8").splitlines()[0].lstrip("# ").strip()
    rows.append((f"scripts/{f.name}", first.split(" - ", 1)[-1]))
TBL(["script", "what it does"], rows, [6.2, 10.2], size=8.5)
TBL(["notebook (review/notebooks/)", "what it shows"], [
    ("01_data_and_taxonomy", "classes, split, imbalance, augmentation"), ("02_model_code_and_training", "model code, parameter counts, training settings, training curves"),
    ("03_results_and_analysis", "accuracy, confusion, calibration, imbalance, classical baselines, pca"), ("04_safety_mechanisms", "uncertainty, conformal routing, cross-check, contamination index, decision engine"),
    ("05_live_demo_federated_video", "full pipeline on real images, grad-cam, federated simulation, video inventory"),
    ("06_colab_training_executed_annotated", "the real Colab training run with explanations (where each feature is coded and why)")], [6.2, 10.2])

# ------------------------------------------------------------------------------------------------ 14
H("14. Ten-minute demo plan for the review")
TBL(["Minute", "Show", "Say"], [
    ("0-1", "slides 1-5 or this dossier section 1-2", "problem and goals in one minute"),
    ("1-3", "notebook 06 (Colab run): scroll the steps, show the training log and the final 93.83 / 97.28 / 98.37", "this is the real run; point to the where-is-it-coded table"),
    ("3-5", "notebook 02: model code, parameter count, attention weights", "two networks, attention decides how much to trust each"),
    ("5-7", "notebook 04: conformal table, cross-check example, contamination index", "safety: stated hazard limit, unsure goes to a person"),
    ("7-9", "notebook 05: eight real images with decisions, grad-cam, video inventory", "the working system"),
    ("9-10", "limits and next step", "hardware prototype with ESP32; sensors are simulated for now")], [2.0, 8.0, 6.4])
P("Commands: all notebooks open from review/notebooks and run from the project folder (they find it by themselves). The live pipeline: python scripts/prototype_sorter.py --mock-node --image <photo>.", italic=True)

# ------------------------------------------------------------------------------------------------ 15
H("15. Questions the professor may ask")
QA = [
    ("Why 33 classes?", "Faculty asked for more than 15 to 20 classes. 33 fine classes in 9 families also let us report recognition and routing separately."),
    ("Which optimiser, activation, loss?", "AdamW (lr 3e-4, weight decay 0.05), GELU in all hidden layers and softmax for outputs, cross-entropy with label smoothing 0.1."),
    ("Is the model overfitting?", "Mildly, not harmfully: train 97.5% vs validation 93.8% in the Colab run; best epoch 8; early stopping after 13. Regularisation: pretraining, augmentation, dropout, weight decay, label smoothing, early stopping."),
    ("How is class imbalance handled?", "Balanced sampler. We found the first run corrected twice (sampler and loss weights), measured the effect (small) and now correct once."),
    ("Did you use PCA, LDA, SVM?", "Yes, as baselines and analysis (section 8, notebook 03). Best classical on hand-made features 71.1%; RBF SVM on frozen deep features 93.3%."),
    ("Why a hybrid of two networks?", "Convolution is good at texture, attention at shape. The learned weights show ConvNeXt gets 82-93%. We did not test ConvNeXt alone, so we do not claim the hybrid is better than it."),
    ("Why transfer learning?", "28 thousand images are too few to train 50 million parameters from scratch; from scratch gave 49% in an earlier try."),
    ("How do you know the test set is clean?", "Stratified fixed split, choices made on validation only, two leaks found and fixed, and the evaluation command checks the validation accuracy stored in the checkpoint."),
    ("Is 93.83% verified?", "It is printed by the executed Colab notebook on its own test split (cell outputs are saved). A local re-score was impossible because the local split differs; it gave a too-high 96.8%, which we discarded."),
    ("What is calibration?", "Among predictions with 80% confidence, about 80% should be right. We measured ECE 16.6% and reduced it to 7.3% with temperature 0.71."),
    ("What is monte carlo dropout?", "Run the head many times with dropout on; the spread of answers is the uncertainty. It is approximate Bayesian inference."),
    ("Is monte carlo dropout good?", "Not better than the plain highest probability at finding mistakes (AUROC 0.60 vs 0.66). The benefit comes from the two-tier gate and a threshold set from a review budget."),
    ("What is conformal prediction?", "A way to output a set of possible answers that contains the truth with a chosen probability, for new data like the calibration data. We use one threshold per family and a stricter one for hazards."),
    ("What does the guarantee cover?", "Items that look like the calibration items. Not very different images; there the uncertainty gate helps but the limit is not kept."),
    ("Why not accuracy only?", "Accuracy hides the cost of a missed battery. We report family accuracy (the bin) and hazard recall separately."),
    ("What is the contamination index?", "A logistic model on two sensor readings that tells if an item is dirty with food. Three models so that one sensor failing does not break it."),
    ("Are the sensors real?", "Not yet. Results use simulated sensor data. The ESP32 prototype is planned and its code is ready but not run on hardware."),
    ("Is federated learning real?", "A simulation with five units on saved features; only weights are exchanged. Not tested on physical devices."),
    ("Edge deployment?", "Models are exported to ONNX (match within 9.4e-6). 14 ms on the laptop GPU and 116 ms on the CPU; Raspberry Pi not measured."),
    ("How is it explained?", "Grad-CAM, LIME, SHAP and the SAM outline. Grad-CAM covers only the ConvNeXt stream."),
    ("What is new in your work?", "Mostly the combination: hazard-exempt cross-check, two-tier hazard gate, conformal routing by material family with a stricter hazard level, and a contamination index that survives a sensor failure. The single techniques (conformal prediction, dropout uncertainty, FedAvg) are published and are not claimed alone."),
    ("What are the limits?", "See section 12: no real sensors, domain gap, conformal limit under shift, no ConvNeXt-only test, simulated federated learning."),
    ("What is Technology Readiness Level 3?", "Proof of concept: the idea works in software on real data. Hardware validation is the next level."),
    ("Where is X coded?", "Section 13 gives a file for every feature; notebook 06 gives the training pipeline step by step."),
    ("What did you do about the detector?", "YOLO26n was fine-tuned on TACO for a 1-class locator (mAP50 0.700) and an 18-class object namer (0.383). The object name feeds the cross-check."),
    ("What happens when the system is unsure about a hazard?", "It goes to priority review, kept out of every recycling bin, never to an automatic gate."),
]
for q, a in QA:
    p = d.add_paragraph(); r = p.add_run("Q: " + q); r.bold = True
    d.add_paragraph("A: " + a)

# ------------------------------------------------------------------------------------------------ 16
H("16. Words explained")
TBL(["Word", "Simple meaning"], [
    ("item / family", "item = exact class (plastic water bottle); family = the bin (plastic)"), ("hazard recall", "of all dangerous items, the share that were called dangerous"),
    ("transfer learning", "start from a network that already learned on millions of pictures"), ("fine-tuning", "keep training that network on our data"),
    ("attention", "a way to give more weight to the more useful part of the input"), ("softmax", "turns scores into probabilities that add up to 1"),
    ("label smoothing", "do not ask for 100% on the true class; makes the model less over-confident"), ("epoch", "one pass over all training images"),
    ("early stopping", "stop training when validation stops improving"), ("calibration / ECE", "do the confidence numbers match the real accuracy / the size of the mismatch"),
    ("temperature scaling", "divide scores by one number to fix over- or under-confidence"), ("monte carlo dropout", "run the model many times with random dropout; spread = uncertainty"),
    ("conformal prediction", "make a set of answers with a chosen chance of containing the truth"), ("leakage", "a hazard sent to a recycling bin (also: test images that were used in training)"),
    ("review workload", "share of items sent to a person"), ("PCA", "reduce many features to a few that keep most information"), ("LDA", "a linear method that finds directions that separate classes"),
    ("SVM", "a classifier that finds the widest margin between classes"), ("bias / variance", "too simple a model (bias) vs too sensitive to the training set (variance)"),
    ("federated learning", "learn together by sharing model weights, not data"), ("ONNX", "a file format to run a model on other devices"), ("TRL", "scale 1 to 9 of how ready a technology is"),
], [4.5, 11.9])

d.save(str(OUT))
print("saved", OUT)
